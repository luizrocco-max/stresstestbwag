"""Saída: resumo no terminal e relatório Excel."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fatores import FATORES
from .motor import Resultado

AZUL = "1F3864"
VERMELHO = "C00000"
CINZA = "F2F2F2"


def _pct(x) -> str:
    return "" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x:+.1%}"


def resumo_terminal(res: Resultado) -> str:
    cart = res.carteira
    linhas = [f"\n== Stress test: {cart.nome} ==", ""]
    linhas.append("Posições:")
    for p in cart.posicoes:
        linhas.append(f"  {p.peso:6.1%}  {p.nome}")
    linhas.append("")
    linhas.append("Carteira por cenário (P&L estimado no período do cenário):")
    tab = res.carteira_cenarios.copy()
    for _, r in tab.sort_values("pl_carteira").iterrows():
        linhas.append(f"  {r['pl_carteira']:+7.1%}  {r['cenario']:<55} pior: {r['pior_posicao'][:35]} ({_pct(r['pl_pior_posicao'])})")
    if not res.betas.empty:
        linhas.append("")
        linhas.append("Betas (excesso de retorno diário x fatores):")
        b = betas_legiveis(res)
        b["nome"] = b["nome"].str[:40]
        linhas.append(b.to_string(index=False, float_format=lambda v: f"{v:6.2f}"))
    if res.avisos:
        linhas.append("")
        linhas.append("Avisos:")
        linhas += [f"  - {a}" for a in res.avisos]
    return "\n".join(linhas)


def betas_legiveis(res: Resultado) -> pd.DataFrame:
    """Betas com fatores de taxa convertidos para '% por +100 bps' (beta x 100)."""
    cols = ["nome", "cnpj"] + res.fatores + ["r2"]
    b = res.betas[cols].copy()
    for f in res.fatores:
        if FATORES[f]["tipo"] == "taxa":
            b[f] = b[f] * 100
            b = b.rename(columns={f: f"{f} (% por +100bps)"})
    return b


def _formatar_aba(ws, pct_cols: set[str], larguras: dict | None = None, congelar: str = "B2"):
    ws.freeze_panes = congelar
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    cab = [c.value for c in ws[1]]
    for j, nome in enumerate(cab, start=1):
        letra = get_column_letter(j)
        largura = (larguras or {}).get(nome, 14 if nome not in ("nome", "cenario", "descricao", "pior_posicao") else 45)
        ws.column_dimensions[letra].width = largura
        if nome in pct_cols and ws.max_row > 1:
            faixa = f"{letra}2:{letra}{ws.max_row}"
            for linha in ws[faixa]:
                for c in linha:
                    c.number_format = "0.0%"
            ws.conditional_formatting.add(faixa, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=VERMELHO)))
    for linha in ws.iter_rows(min_row=2):
        for c in linha:
            if isinstance(c.value, (datetime, date)):
                c.number_format = "dd/mm/yyyy"
    ws.row_dimensions[1].height = 32


def excel(res: Resultado, caminho: str | Path) -> Path:
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    fat = res.fatores
    taxa_cols = {f for f in fat if FATORES[f]["tipo"] == "taxa"}
    ret_cols = {f for f in fat if FATORES[f]["tipo"] == "retorno"}

    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        # Resumo
        resumo = res.carteira_cenarios.reset_index(drop=True)[
            ["cenario", "tipo", "n_dias", "pl_carteira", "pl_modelo", "cdi_periodo", "peso_com_historico_real",
             "fatores_sem_dado", "pior_posicao", "pl_pior_posicao", "descricao"]]
        resumo.to_excel(xw, sheet_name="Resumo", index=False)
        _formatar_aba(xw.sheets["Resumo"], {"pl_carteira", "pl_modelo", "cdi_periodo", "peso_com_historico_real", "pl_pior_posicao"})

        # Matriz fundo x cenário
        mat = res.matriz.copy()
        pesos = {p.nome: p.peso for p in res.carteira.posicoes}
        mat.insert(0, "peso", [pesos.get(n, np.nan) for n in mat.index])
        mat.loc["CARTEIRA"] = [1.0] + list(res.carteira_cenarios.set_index("cenario").loc[mat.columns[1:], "pl_carteira"])
        mat = mat.reset_index().rename(columns={"index": "nome"})
        mat.to_excel(xw, sheet_name="Fundos x Cenarios", index=False)
        _formatar_aba(xw.sheets["Fundos x Cenarios"], set(mat.columns[1:]), larguras={c: 18 for c in mat.columns[2:]})

        # Detalhe
        det = res.detalhe[["nome", "cnpj", "peso", "cenario", "tipo_cenario", "fonte", "usado", "historico",
                           "modelo_total", "modelo_excesso", "cdi"]]
        det.to_excel(xw, sheet_name="Detalhe", index=False)
        _formatar_aba(xw.sheets["Detalhe"], {"peso", "usado", "historico", "modelo_total", "modelo_excesso", "cdi"})

        # Betas
        if not res.betas.empty:
            b = betas_legiveis(res).drop(columns=["r2"])
            extras = res.betas[["r2", "r2_ajustado", "alpha_anual", "vol_residual_anual", "n_obs",
                                "janela_inicio", "janela_fim"] + [f"t_{f}" for f in fat]]
            b = pd.concat([b, extras], axis=1)
            b.to_excel(xw, sheet_name="Betas", index=False)
            _formatar_aba(xw.sheets["Betas"], {"r2", "r2_ajustado", "alpha_anual", "vol_residual_anual"},
                          larguras={c: 16 for c in b.columns if "(%" in c})
            ws = xw.sheets["Betas"]
            for j, nome in enumerate([c.value for c in ws[1]], start=1):
                if nome in fat or "(% por" in nome or nome.startswith("t_"):
                    for linha in ws[f"{get_column_letter(j)}2:{get_column_letter(j)}{ws.max_row}"]:
                        for c in linha:
                            c.number_format = "0.00"

        # Estatísticas
        if not res.estatisticas.empty:
            st = res.estatisticas.copy()
            for c in st.columns:
                if c.startswith("data_") or c.endswith("_serie"):
                    st[c] = pd.to_datetime(st[c]).dt.date
            st.to_excel(xw, sheet_name="Estatisticas", index=False)
            _formatar_aba(xw.sheets["Estatisticas"], {"vol_anual", "max_drawdown", "pior_dia", "pior_21d", "var95_21d",
                                                     "es95_21d", "ret_12m", "ret_36m"})

        # Cenários (choques)
        ch = res.choques.reset_index()[["id", "cenario", "tipo", "inicio", "fim", "n_dias"] + fat + ["CDI", "descricao"]]
        ch.to_excel(xw, sheet_name="Cenarios", index=False)
        _formatar_aba(xw.sheets["Cenarios"], ret_cols | {"CDI"})
        ws = xw.sheets["Cenarios"]
        for j, nome in enumerate([c.value for c in ws[1]], start=1):
            if nome in taxa_cols:
                for linha in ws[f"{get_column_letter(j)}2:{get_column_letter(j)}{ws.max_row}"]:
                    for c in linha:
                        c.number_format = '+0.00" p.p.";-0.00" p.p."'

        # Contribuições
        contrib = res.contribuicoes.reset_index(drop=True)
        contrib["total_modelo"] = res.carteira_cenarios["pl_modelo"].values
        contrib.to_excel(xw, sheet_name="Contribuicoes", index=False)
        _formatar_aba(xw.sheets["Contribuicoes"], set(fat) | {"CDI", "total_modelo"})

        # Fatores (legenda) e avisos
        leg = pd.DataFrame([{"fator": k, "descricao": v["nome"], "tipo": v["tipo"], "unidade_do_choque": v["unidade"],
                             "leitura_do_beta": ("retorno do fundo por 1% do fator (0,5 = metade do índice)" if v["tipo"] == "retorno"
                                                 else "retorno do fundo (%) por +100 bps na taxa (-2 = cai 2% se o juro sobe 1 p.p.)"),
                             "usado": k in fat} for k, v in FATORES.items()])
        leg.to_excel(xw, sheet_name="Fatores", index=False)
        _formatar_aba(xw.sheets["Fatores"], set(), larguras={"descricao": 45, "leitura_do_beta": 80})
        av = pd.DataFrame({"aviso": res.avisos or ["nenhum"]})
        av.loc[len(av)] = f"gerado em {datetime.now():%d/%m/%Y %H:%M}"
        av.to_excel(xw, sheet_name="Avisos", index=False)
        _formatar_aba(xw.sheets["Avisos"], set(), larguras={"aviso": 120})
    return caminho
